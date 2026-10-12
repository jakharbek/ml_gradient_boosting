# -*- coding: utf-8 -*-
"""Excel-версия дополнительного урока 01 «Бустинг с нуля на четырёх ноутбуках».

Запуск: python extra/01_boosting_by_hand/build_excel.py  →  boosting_by_hand.xlsx рядом со скриптом.
Один лист, всё формулами: исходные данные (жёлтые ячейки), шаг 0 (среднее), шаг 1 (экран),
шаг 2 (SSD), финальная формула и сводка ошибки 600 → 120 → 0.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as cl

wb = Workbook()
ws = wb.active
ws.title = "Градиентный бустинг"
ws.sheet_view.showGridLines = False

F = "Arial"
NAVY, TEAL, GREY, LIGHT, PALE = "1F3864", "2E75B6", "595959", "DEEAF6", "F2F6FB"
GREEN_F, GREEN_B = "006100", "E2F0D9"
ORANGE_B, ORANGE_F = "FFF2CC", "7F6000"
WHITE = "FFFFFF"

thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color=NAVY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", indent=1)
wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

def fill(c): return PatternFill("solid", fgColor=c)

def cell(r, c, v=None, *, font=None, bg=None, border=box, align=center, fmt=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = font or Font(name=F, size=11)
    if bg: x.fill = fill(bg)
    if border: x.border = border
    if align: x.alignment = align
    if fmt: x.number_format = fmt
    return x

NUM = '#,##0;[Red]-#,##0;0'

def banner(r, text, sub=None):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    cell(r, 2, text, font=Font(name=F, size=13, bold=True, color=WHITE), bg=NAVY, border=None, align=left)
    ws.row_dimensions[r].height = 26
    if sub:
        ws.merge_cells(start_row=r+1, start_column=2, end_row=r+1, end_column=8)
        cell(r+1, 2, sub, font=Font(name=F, size=10, italic=True, color=GREY), bg=PALE, border=None, align=wrap)
        ws.row_dimensions[r+1].height = 30
        return r + 3
    return r + 2

def header(r, labels, c0=2):
    for i, lab in enumerate(labels):
        cell(r, c0+i, lab, font=Font(name=F, size=11, bold=True, color=WHITE), bg=TEAL)
    ws.row_dimensions[r].height = 22

def label(r, text, c0=2, span=3):
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0+span-1)
    cell(r, c0, text, font=Font(name=F, size=11, bold=True, color=NAVY), bg=LIGHT, align=left)

def note(r, c, text):
    cell(r, c, text, font=Font(name=F, size=10, italic=True, color=GREY), border=None, align=left)

def result(r, c, v):
    return cell(r, c, v, font=Font(name=F, size=11, bold=True, color=GREEN_F), bg=GREEN_B, fmt=NUM)

def inp(r, c, v):
    return cell(r, c, v, font=Font(name=F, size=11, color="0000FF"), bg=ORANGE_B, fmt=NUM)

# ---------- Title ----------
ws.merge_cells("B2:H2")
cell(2, 2, "Градиентный бустинг — пример на 4 ноутбуках", font=Font(name=F, size=18, bold=True, color=NAVY), border=None, align=left)
ws.row_dimensions[2].height = 34
ws.merge_cells("B3:H3")
cell(3, 2, "Каждая следующая модель исправляет ошибки суммы предыдущих. Жёлтые ячейки — исходные данные, их можно менять; всё остальное пересчитается само.",
     font=Font(name=F, size=10, italic=True, color=GREY), border=None, align=wrap)
ws.row_dimensions[3].height = 30

r = 5
# ---------- Data ----------
r = banner(r, "Исходные данные")
header(r, ["Ноутбук", "Экран", "SSD", "Реальная цена"]); r += 1
data = [("A", '13"', "256 ГБ", 820), ("B", '13"', "512 ГБ", 880),
        ("C", '15"', "256 ГБ", 1120), ("D", '15"', "512 ГБ", 1180)]
d0 = r
for n, s, ssd, p in data:
    cell(r, 2, n, font=Font(name=F, bold=True)); cell(r, 3, s); cell(r, 4, ssd); inp(r, 5, p); r += 1
d1 = r - 1
PR = f"E{d0}:E{d1}"
r += 1

# ---------- Step 0 ----------
r = banner(r, "Шаг 0 — стартовое предсказание",
           "Пока модели нет, предсказываем всем одно число — среднее всех цен. Остаток = реальная − предсказание: насколько промахнулись.")
label(r, "Сумма цен"); cell(r, 5, f"=SUM({PR})", fmt=NUM); note(r, 6, "820 + 880 + 1120 + 1180"); r += 1
label(r, "Количество"); cell(r, 5, f"=COUNT({PR})"); r += 1
label(r, "Среднее = стартовое предсказание"); mean = result(r, 5, f"=E{r-2}/E{r-1}"); MEAN = f"$E${r}"; note(r, 6, "сумма / количество"); r += 2
header(r, ["Ноутбук", "Реальная", "Предсказание", "Остаток", "|Остаток|"]); r += 1
s0 = r
for i, (n, _, _, p) in enumerate(data):
    cell(r, 2, n, font=Font(name=F, bold=True))
    cell(r, 3, f"=E{d0+i}", fmt=NUM)
    cell(r, 4, f"={MEAN}", fmt=NUM)
    result(r, 5, f"=C{r}-D{r}")
    cell(r, 6, f"=ABS(E{r})", fmt=NUM)
    r += 1
RES0 = f"E{s0}:E{r-1}"
label(r, "Сумма модулей ошибок", span=4); result(r, 6, f"=SUM(F{s0}:F{r-1})"); SUM0 = f"F{r}"; r += 2

# ---------- Step 1 ----------
r = banner(r, "Шаг 1 — модель по экрану",
           "Делим на группы по экрану. Поправка группы = среднее остатков в ней. Новое предсказание = старое + поправка.")
header(r, ["Группа", "Состав", "Сумма остатков", "Кол-во", "Поправка"]); r += 1
cell(r, 2, '13"', font=Font(name=F, bold=True)); cell(r, 3, "A, B"); cell(r, 4, f"=E{s0}+E{s0+1}", fmt=NUM); cell(r, 5, 2); result(r, 6, f"=D{r}/E{r}"); P13 = f"$F${r}"; r += 1
cell(r, 2, '15"', font=Font(name=F, bold=True)); cell(r, 3, "C, D"); cell(r, 4, f"=E{s0+2}+E{s0+3}", fmt=NUM); cell(r, 5, 2); result(r, 6, f"=D{r}/E{r}"); P15 = f"$F${r}"; r += 2
header(r, ["Ноутбук", "Старое", "Поправка", "Новое", "Реальная", "Новый остаток", "|Остаток|"]); r += 1
s1 = r
for i, (n, _, _, p) in enumerate(data):
    cell(r, 2, n, font=Font(name=F, bold=True))
    cell(r, 3, f"=D{s0+i}", fmt=NUM)
    cell(r, 4, f"={P13 if i < 2 else P15}", fmt=NUM)
    result(r, 5, f"=C{r}+D{r}")
    cell(r, 6, f"=E{d0+i}", fmt=NUM)
    result(r, 7, f"=F{r}-E{r}")
    cell(r, 8, f"=ABS(G{r})", fmt=NUM)
    r += 1
label(r, "Сумма модулей ошибок", span=6); result(r, 8, f"=SUM(H{s1}:H{r-1})"); SUM1 = f"H{r}"; r += 2

# ---------- Step 2 ----------
r = banner(r, "Шаг 2 — модель по SSD",
           "Остатки после шага 1 — это то, чего экран не объяснил. Теперь делим по SSD и повторяем ту же процедуру.")
header(r, ["Группа", "Состав", "Сумма остатков", "Кол-во", "Поправка"]); r += 1
cell(r, 2, "256 ГБ", font=Font(name=F, bold=True)); cell(r, 3, "A, C"); cell(r, 4, f"=G{s1}+G{s1+2}", fmt=NUM); cell(r, 5, 2); result(r, 6, f"=D{r}/E{r}"); P256 = f"$F${r}"; r += 1
cell(r, 2, "512 ГБ", font=Font(name=F, bold=True)); cell(r, 3, "B, D"); cell(r, 4, f"=G{s1+1}+G{s1+3}", fmt=NUM); cell(r, 5, 2); result(r, 6, f"=D{r}/E{r}"); P512 = f"$F${r}"; r += 2
header(r, ["Ноутбук", "Старое", "Поправка", "Новое", "Реальная", "Новый остаток", "|Остаток|"]); r += 1
s2 = r
for i, (n, _, _, p) in enumerate(data):
    cell(r, 2, n, font=Font(name=F, bold=True))
    cell(r, 3, f"=E{s1+i}", fmt=NUM)
    cell(r, 4, f"={P256 if i % 2 == 0 else P512}", fmt=NUM)
    result(r, 5, f"=C{r}+D{r}")
    cell(r, 6, f"=E{d0+i}", fmt=NUM)
    result(r, 7, f"=F{r}-E{r}")
    cell(r, 8, f"=ABS(G{r})", fmt=NUM)
    r += 1
label(r, "Сумма модулей ошибок", span=6); result(r, 8, f"=SUM(H{s2}:H{r-1})"); SUM2 = f"H{r}"; r += 2

# ---------- Step 3 ----------
r = banner(r, "Шаг 3 — финальная формула",
           "Итог = старт + поправка модели 1 + поправка модели 2. Каждая модель выучила свой кусок, сумма дала точный ответ.")
header(r, ["Ноутбук", "Старт", "Экран", "SSD", "Итог", "Реальная", "Ошибка"]); r += 1
s3 = r
for i, (n, _, _, p) in enumerate(data):
    cell(r, 2, n, font=Font(name=F, bold=True))
    cell(r, 3, f"={MEAN}", fmt=NUM)
    cell(r, 4, f"={P13 if i < 2 else P15}", fmt=NUM)
    cell(r, 5, f"={P256 if i % 2 == 0 else P512}", fmt=NUM)
    result(r, 6, f"=C{r}+D{r}+E{r}")
    cell(r, 7, f"=E{d0+i}", fmt=NUM)
    cell(r, 8, f"=G{r}-F{r}", fmt=NUM)
    r += 1
r += 1

# ---------- Summary ----------
r = banner(r, "Как уменьшалась ошибка")
header(r, ["Шаг", "Что сделали", "Сумма модулей ошибок"]); r += 1
for lab, what, ref in [("0", "Среднее для всех", SUM0), ("1", "+ модель по экрану", SUM1), ("2", "+ модель по SSD", SUM2)]:
    cell(r, 2, lab, font=Font(name=F, bold=True)); cell(r, 3, what, align=left); result(r, 4, f"={ref}"); r += 1
r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
cell(r, 2, "Главная идея: бустинг не строит одну умную модель, а складывает цепочку простых, где каждая следующая исправляет ошибки суммы предыдущих.",
     font=Font(name=F, size=11, bold=True, color=NAVY), bg=LIGHT, border=None, align=wrap)
ws.row_dimensions[r].height = 36

for c, w in {1: 3, 2: 30, 3: 18, 4: 18, 5: 18, 6: 18, 7: 18, 8: 16, 9: 3}.items():
    ws.column_dimensions[cl(c)].width = w

wb.save(Path(__file__).with_name("boosting_by_hand.xlsx"))
print("ok")
