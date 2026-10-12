# -*- coding: utf-8 -*-
"""Excel-пособие: градиентный бустинг с learning rate, 20 шагов.
Листы: Содержание · Теория · Сводка (графики) · Шаги (20 сворачиваемых блоков A–D с живыми пояснениями)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.formatting.rule import DataBarRule, CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as cl
from openpyxl.worksheet.hyperlink import Hyperlink

F = "Arial"
NAVY, TEAL, GREY, LIGHT, PALE = "1F3864", "2E75B6", "595959", "DEEAF6", "F2F6FB"
GREEN_F, GREEN_B, ORANGE_B, ORANGE_F = "006100", "E2F0D9", "FCE4D6", "C55A11"
INPUT_B, INPUT_F, EXPL_B, EXPL_LINE, WHITE = "FFF2CC", "0000FF", "FBF7EE", "E8DCC0", "FFFFFF"
thin = Side(style="thin", color="BFBFBF"); med = Side(style="medium", color=NAVY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", indent=1)
wrap = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
NUM, INT = '#,##0.00;[Red]-#,##0.00;0', '#,##0;[Red]-#,##0;0'
def fill(c): return PatternFill("solid", fgColor=c)

wb = Workbook()

# ============================================================
# generic helpers bound to a sheet
# ============================================================
class S:
    def __init__(self, ws, last=13):
        self.ws, self.last = ws, last
        ws.sheet_view.showGridLines = False
    def cell(self, r, c, v=None, *, font=None, bg=None, border=box, align=center, fmt=NUM):
        x = self.ws.cell(row=r, column=c, value=v)
        x.font = font or Font(name=F, size=10)
        if bg: x.fill = fill(bg)
        if border: x.border = border
        if align: x.alignment = align
        if fmt: x.number_format = fmt
        return x
    def merge(self, r, c1, c2, v, **kw):
        self.ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        return self.cell(r, c1, v, **kw)
    def banner(self, r, text, sub=None, c1=2):
        self.merge(r, c1, self.last, text, font=Font(name=F, size=12, bold=True, color=WHITE), bg=NAVY, border=None, align=left, fmt=None)
        self.ws.row_dimensions[r].height = 24
        if sub:
            self.merge(r+1, c1, self.last, sub, font=Font(name=F, size=9, italic=True, color=GREY), bg=PALE, border=None, align=wrap, fmt=None)
            self.ws.row_dimensions[r+1].height = 30
            return r + 3
        return r + 2
    def header(self, r, labels, c0=2):
        for i, lab in enumerate(labels):
            self.cell(r, c0+i, lab, font=Font(name=F, size=9, bold=True, color=WHITE), bg=TEAL,
                      align=Alignment(horizontal="center", vertical="center", wrap_text=True), fmt=None)
        self.ws.row_dimensions[r].height = 30
    def label(self, r, text, c0=2, span=2):
        return self.merge(r, c0, c0+span-1, text, font=Font(name=F, size=10, bold=True, color=NAVY), bg=LIGHT, align=left, fmt=None)
    def explain(self, r, formula, height=46, c1=2):
        self.merge(r, c1, self.last, formula, font=Font(name=F, size=9, color="3B3B3B"), bg=EXPL_B,
                   border=Border(left=Side(style="thin", color=EXPL_LINE), right=Side(style="thin", color=EXPL_LINE),
                                 top=Side(style="thin", color=EXPL_LINE), bottom=Side(style="thin", color=EXPL_LINE)), align=wrap, fmt=None)
        self.ws.row_dimensions[r].height = height
        return r + 1
    def res(self, r, c, v, fmt=NUM):
        return self.cell(r, c, v, font=Font(name=F, size=10, bold=True, color=GREEN_F), bg=GREEN_B, fmt=fmt)
    def inp(self, r, c, v, fmt=INT):
        return self.cell(r, c, v, font=Font(name=F, size=10, bold=True, color=INPUT_F), bg=INPUT_B, fmt=fmt)
    def link(self, r, c, text, target, font=None):
        x = self.cell(r, c, text, font=font or Font(name=F, size=10, color=TEAL, underline="single"), border=None, align=left, fmt=None)
        x.hyperlink = Hyperlink(ref=x.coordinate, location=target.lstrip("#"), display=str(text))
        return x

def T(ref): return f'TEXT({ref},"0.##")'

# ============================================================
# Sheet: Содержание
# ============================================================
ws0 = wb.active; ws0.title = "Содержание"; s0 = S(ws0, last=8)
ws0.column_dimensions["A"].width = 2; ws0.column_dimensions["B"].width = 34; ws0.column_dimensions["C"].width = 70
s0.cell(2, 2, "Градиентный бустинг с learning rate — 20 шагов", font=Font(name=F, size=18, bold=True, color=NAVY), border=None, align=left, fmt=None)
ws0.row_dimensions[2].height = 34
s0.merge(3, 2, 8, "Учебное пособие в Excel: один пример (четыре ноутбука), один параметр ν, двадцать шагов — и у каждого числа есть формула, обоснование и живой пересчёт. Жёлтые ячейки — входы; всё остальное считается формулами.",
         font=Font(name=F, size=10, italic=True, color=GREY), border=None, align=wrap, fmt=None)
ws0.row_dimensions[3].height = 34
toc_rows = []  # filled later: (text, target, desc)
TOC_START = 6

# ============================================================
# Sheet: Теория
# ============================================================
th = wb.create_sheet("Теория"); st = S(th, last=2)
th.column_dimensions["A"].width = 2; th.column_dimensions["B"].width = 118
def TT(r, text, kind="p"):
    c = th.cell(row=r, column=2, value=text)
    if kind == "h1": c.font = Font(name=F, size=16, bold=True, color=NAVY); th.row_dimensions[r].height = 30
    elif kind == "h2": c.font = Font(name=F, size=12, bold=True, color=WHITE); c.fill = fill(NAVY); c.alignment = left; th.row_dimensions[r].height = 24
    elif kind == "f": c.font = Font(name="Cambria Math", size=12, bold=True, color=NAVY); c.fill = fill(LIGHT); c.alignment = center; th.row_dimensions[r].height = 26
    else: c.font = Font(name=F, size=10); c.alignment = wrap; th.row_dimensions[r].height = 15 * max(1, len(text)//112 + 1)
    return r + 1
r = 2
r = TT(r, "Математическое основание каждого расчёта", "h1")
st.link(r, 2, "← Содержание", "#'Содержание'!A1"); r += 2
TH = {}
sections = [
 ("1. Постановка задачи", [
   ("p", "Есть n = 4 объекта (ноутбуки A, B, C, D) с признаками x (экран, SSD) и целевой переменной y (реальная цена). Строим функцию F(x), прогноз ŷ = F(x), и хотим, чтобы прогноз был как можно ближе к y на всех объектах сразу."),
   ("p", "«Как можно ближе» измеряет функция потерь. Здесь — квадратичная: L(y, F) = ½·(y − F)². Она гладкая (есть производная), штрафует большие промахи сильнее малых, и у её минимума простой ответ — среднее (раздел 2).")]),
 ("2. Шаг 0: почему стартовое предсказание — среднее", [
   ("p", "Пока признаков не используем, модель — одно число c для всех. Минимизируем сумму квадратов ошибок:"),
   ("f", "S(c) = Σᵢ (yᵢ − c)²    →    dS/dc = −2·Σᵢ (yᵢ − c) = 0    →    c = (1/n)·Σᵢ yᵢ = ȳ"),
   ("p", "Производная равна нулю ровно в среднем арифметическом: F₀ = ȳ = (820 + 880 + 1120 + 1180) / 4 = 1000. Это решение задачи минимизации, а не договорённость.")]),
 ("3. Остаток и его связь с градиентом", [
   ("p", "Остаток (residual) объекта i на шаге m: rᵢ = yᵢ − F_{m−1}(xᵢ) — насколько текущая модель промахнулась. Минус — завысили, плюс — занизили."),
   ("f", "rᵢ = yᵢ − F_{m−1}(xᵢ) = −∂L(yᵢ, F)/∂F  при  L = ½(y − F)²"),
   ("p", "Для квадратичной потери остаток — в точности антиградиент потери по прогнозу. Значит «учить следующую модель на остатках» = «делать шаг градиентного спуска в пространстве функций». Отсюда слово «градиентный». Для других потерь вместо остатка берут антиградиент — формула меняется, идея нет.")]),
 ("4. Этап B: группировка по колонкам и поправка группы", [
   ("p", "Слабая модель h_m — пень (stump): одна колонка делит объекты на две группы, каждая группа получает одно число — поправку γ. Какое γ лучше всего для группы G?"),
   ("f", "γ_G = argmin_γ Σ_{i∈G} (rᵢ − γ)²  =  (1/|G|)·Σ_{i∈G} rᵢ  — среднее остатков группы"),
   ("p", "Та же выкладка, что в разделе 2, но по объектам группы. Группируем по колонке-признаку, а не по остаткам: признак известен у любого нового объекта, остаток — нет. Столбец «Квадраты после поправки» на листе «Шаги» — это Σ_{i∈G}(rᵢ − γ_G)²: ошибка, которая останется в группе после поправки.")]),
 ("5. Этап C: критерий выбора колонки", [
   ("f", "j* = argmin_j  [ Σ_{i∈G₁(j)} (rᵢ − γ₁)² + Σ_{i∈G₂(j)} (rᵢ − γ₂)² ]"),
   ("p", "Для каждой колонки складываем ошибку, оставшуюся в её двух группах, и берём минимум. Эквивалентно: колонка с наибольшим уменьшением суммы квадратов (gain). На шаге 1 экран оставляет 3 600, SSD — 90 000: экран объясняет в 25 раз больше. Так дерево «само находит» важный признак.")]),
 ("6. Этап D: learning rate и обновление модели", [
   ("f", "F_m(x) = F_{m−1}(x) + ν · h_m(x),   0 < ν ≤ 1"),
   ("p", "ν (ню) — темп обучения (learning rate). При ν = 1 прибавляется вся поправка, при ν = 0.5 — половина. Это регуляризация по Фридману: каждое дерево делает заведомо неполный шаг, итоговая модель — сумма многих осторожных поправок. На практике ν ∈ [0.01; 0.3], а число деревьев подбирается под него."),
   ("p", "Почему не ν = 1, если быстрее? На обучающих данных — да. Но цель — новые объекты: резкие шаги подстраиваются под шум обучающей выборки, маленькие — усредняют его. Меньший ν при большем числе деревьев почти всегда даёт лучшее качество на новых данных.")]),
 ("7. Сходимость: почему ошибка стремится к нулю, но не достигает его", [
   ("p", "Если колонка выбирается снова, её невыученный вклад умножается на (1 − ν): 150 → 75 → 37.5 → 18.75 …"),
   ("f", "остаток вклада после k шагов = (1 − ν)ᵏ · (исходный вклад)  →  0  при  k → ∞,  но ≠ 0 ни при каком k"),
   ("p", "Поэтому на листе «Сводка» Σ|rᵢ| падает 600 → 300 → 150 → … → 0.29 на 20-м шаге и будет падать дальше бесконечно. В реальной задаче это достоинство: модель можно остановить в любой момент, и число шагов — ещё один рычаг против переобучения.")]),
 ("8. Когда остановиться", [
   ("p", "(1) Фиксированное число деревьев n_estimators; (2) остановка, когда улучшение за шаг меньше порога; (3) ранняя остановка (early stopping): ошибка считается на отложенных данных, и обучение прекращается, когда она перестаёт падать. Третье — стандарт практики: оно отвечает не «как выучены эти 4 ноутбука», а «как будет предсказан пятый».")]),
 ("9. Обозначения на листе «Шаги»", [
   ("p", "yᵢ — реальная цена (жёлтые ячейки).  F_{m−1}(xᵢ) — «Старое предсказание».  rᵢ — «Остаток».  γ_G — «Поправка (среднее)».  h_m(xᵢ) — «Поправка выбранной колонки».  ν — «Learning rate».  F_m(xᵢ) — «Новое предсказание».  Σ|rᵢ| — наглядный индикатор; Σrᵢ² — то, что модель минимизирует на самом деле.")]),
]
for title, items in sections:
    TH[title] = r
    r = TT(r, title, "h2")
    for kind, text in items: r = TT(r, text, kind)
    r += 1

# ============================================================
# Sheet: Шаги (built first to know references), then Сводка
# ============================================================
ws = wb.create_sheet("Шаги"); s = S(ws, last=13)
sv = wb.create_sheet("Сводка"); ss = S(sv, last=9)

# ---- Шаги: header & inputs ----
s.merge(2, 2, 13, "Шаги бустинга 1–20: этапы A → B → C → D с живыми пояснениями", font=Font(name=F, size=15, bold=True, color=NAVY), border=None, align=left, fmt=None)
ws.row_dimensions[2].height = 28
s.link(3, 2, "← Содержание", "#'Содержание'!A1"); s.link(3, 3, "Сводка и графики →", "#'Сводка'!A1"); s.link(3, 4, "Теория →", "#'Теория'!A1")
s.merge(4, 2, 13, "Слева у каждого блока — кнопка «−/+» (группировка строк Excel): шаг можно свернуть. Числа в пояснениях — формулы, они пересчитываются при смене ν или цен.",
        font=Font(name=F, size=9, italic=True, color=GREY), border=None, align=wrap, fmt=None)
ws.row_dimensions[4].height = 28
r = 6
r = s.banner(r, "Параметры и исходные данные")
s.label(r, "Learning rate ν"); s.inp(r, 4, 0.5, fmt="0.00"); LR = f"Шаги!$D${r}"; LRc = f"$D${r}"
dv = DataValidation(type="decimal", operator="between", formula1="0.01", formula2="1", showErrorMessage=True,
                    errorTitle="Недопустимое ν", error="Learning rate должен быть в интервале от 0.01 до 1.")
ws.add_data_validation(dv); dv.add(f"D{r}")
s.cell(r, 5, "0 < ν ≤ 1. Поставьте 1 — получите исходный пример за 2 шага; 0.1 — очень медленно.", font=Font(name=F, size=9, italic=True, color=GREY), border=None, align=left, fmt=None)
wb.defined_names["LR"] = DefinedName("LR", attr_text=f"Шаги!$D${r}")
r += 2
s.header(r, ["Ноутбук", "Экран", "SSD", "Реальная цена yᵢ"]); r += 1
data = [("A", '13"', "256 ГБ", 820), ("B", '13"', "512 ГБ", 880), ("C", '15"', "256 ГБ", 1120), ("D", '15"', "512 ГБ", 1180)]
d0 = r
for n, scr, ssd, p in data:
    s.cell(r, 2, n, font=Font(name=F, bold=True), fmt=None); s.cell(r, 3, scr, fmt=None); s.cell(r, 4, ssd, fmt=None); s.inp(r, 5, p); r += 1
REAL = [f"$E${d0+i}" for i in range(4)]
r += 1
s.label(r, "Стартовое предсказание F₀ = ȳ", span=3); s.res(r, 5, f"=AVERAGE(E{d0}:E{d0+3})"); MEAN = f"$E${r}"
wb.defined_names["F0"] = DefinedName("F0", attr_text=f"Шаги!$E${r}")
r = s.explain(r+1, f'="F₀ = ("&{REAL[0]}&" + "&{REAL[1]}&" + "&{REAL[2]}&" + "&{REAL[3]}&") / 4 = "&{T(MEAN)}&".  Среднее — единственное число, минимизирующее Σ(yᵢ − c)²: dS/dc = −2·Σ(yᵢ − c) = 0 ⇒ c = ȳ (Теория, раздел 2)."', 32)
r += 1

N = 20
COLS_D = ["Ноутбук", "F_{m−1}(xᵢ)", "h_m(xᵢ) — поправка выбранной колонки", "ν·h_m(xᵢ)", "F_m(xᵢ) = F_{m−1} + ν·h_m", "yᵢ", "Новый остаток yᵢ − F_m(xᵢ)", "|остаток|", "остаток²"]
groups = [("Экран", '13"', "A, B", [0, 1]), ("Экран", '15"', "C, D", [2, 3]), ("SSD", "256 ГБ", "A, C", [0, 2]), ("SSD", "512 ГБ", "B, D", [1, 3])]
prev_new = [MEAN]*4
step_info = []
step_rows = []

for m in range(1, N+1):
    top = r
    step_rows.append(top)
    s.merge(r, 2, 13, f"Шаг {m}:  F_{m} = F_{m-1} + ν·h_{m}", font=Font(name=F, size=11, bold=True, color=WHITE), bg=TEAL, align=left, fmt=None)
    ws.row_dimensions[r].height = 22
    r += 1
    # live one-line status at top of the step, filled after D is known
    status_row = r; r += 1

    # ---- A ----
    s.label(r, f"A. Остатки текущей модели F_{m-1}   (Теория, раздел 3)", span=5); r += 1
    ra_expl = r; r += 1
    s.header(r, ["Ноутбук", f"Старое предсказание F_{m-1}(xᵢ)", "Реальная yᵢ", f"Остаток rᵢ = yᵢ − F_{m-1}(xᵢ)", "|rᵢ|"]); r += 1
    ra = r
    for i, (n, _, _, _) in enumerate(data):
        s.cell(r, 2, n, font=Font(name=F, bold=True), fmt=None); s.cell(r, 3, f"={prev_new[i]}"); s.cell(r, 4, f"={REAL[i]}", fmt=INT)
        s.res(r, 5, f"=D{r}-C{r}"); s.cell(r, 6, f"=ABS(E{r})"); r += 1
    RESID = [f"E{ra+i}" for i in range(4)]
    s.label(r, "Σ|rᵢ|  и  Σrᵢ²", span=4); s.res(r, 6, f"=SUM(F{ra}:F{ra+3})"); s.res(r, 7, f"=SUMSQ(E{ra}:E{ra+3})"); SA0, SS0 = f"F{r}", f"G{r}"
    ws.conditional_formatting.add(f"F{ra}:F{ra+3}", DataBarRule(start_type="num", start_value=0, end_type="max", color="9DC3E6"))
    s.explain(ra_expl, f'="Остаток — антиградиент квадратичной потери: rᵢ = yᵢ − F_{m-1}(xᵢ) = −∂L/∂F — «что модель ещё не выучила». Для A: "&{T(f"D{ra}")}&" − "&{T(f"C{ra}")}&" = "&{T(f"E{ra}")}&"; для B: "&{T(f"D{ra+1}")}&" − "&{T(f"C{ra+1}")}&" = "&{T(f"E{ra+1}")}&"; для C: "&{T(f"D{ra+2}")}&" − "&{T(f"C{ra+2}")}&" = "&{T(f"E{ra+2}")}&"; для D: "&{T(f"D{ra+3}")}&" − "&{T(f"C{ra+3}")}&" = "&{T(f"E{ra+3}")}&". Минус — завысили, плюс — занизили. Σ|rᵢ| = "&{T(SA0)}&", Σrᵢ² = "&{T(SS0)}&"."', 46)
    r += 2

    # ---- B ----
    s.label(r, "B. Группировка по колонкам и поправки групп   (Теория, раздел 4)", span=9); r += 1
    rb_expl = r; r += 1
    s.header(r, ["Колонка", "Группа G", "Состав", "Σ остатков группы", "|G|", "Поправка γ_G = Σrᵢ / |G|", "Остаётся Σ(rᵢ − γ_G)²"]); r += 1
    rb = r
    for col, gname, members, idx in groups:
        s.cell(r, 2, col, font=Font(name=F, bold=True), fmt=None); s.cell(r, 3, gname, fmt=None); s.cell(r, 4, members, fmt=None)
        s.cell(r, 5, f"={RESID[idx[0]]}+{RESID[idx[1]]}"); s.cell(r, 6, 2, fmt="0"); s.res(r, 7, f"=E{r}/F{r}")
        s.cell(r, 8, f"=({RESID[idx[0]]}-G{r})^2+({RESID[idx[1]]}-G{r})^2"); r += 1
    SCR = {0: f"G{rb}", 1: f"G{rb}", 2: f"G{rb+1}", 3: f"G{rb+1}"}; SSD = {0: f"G{rb+2}", 2: f"G{rb+2}", 1: f"G{rb+3}", 3: f"G{rb+3}"}
    s.explain(rb_expl, f'="Группируем по колонке-признаку (он известен у любого нового объекта), а не по остаткам. Лучшее одно число для группы при квадратичной потере — среднее остатков: γ_G = argmin Σ(rᵢ − γ)² = Σrᵢ / |G|. "&"Экран 13″ {{A, B}}: γ = ("&{T(RESID[0])}&" + "&{T(RESID[1])}&") / 2 = "&{T(f"G{rb}")}&", остаётся "&{T(f"H{rb}")}&";  экран 15″ {{C, D}}: γ = "&{T(f"G{rb+1}")}&", остаётся "&{T(f"H{rb+1}")}&";  SSD 256 {{A, C}}: γ = ("&{T(RESID[0])}&" + "&{T(RESID[2])}&") / 2 = "&{T(f"G{rb+2}")}&", остаётся "&{T(f"H{rb+2}")}&";  SSD 512 {{B, D}}: γ = "&{T(f"G{rb+3}")}&", остаётся "&{T(f"H{rb+3}")}&"."', 58)
    r += 1

    # ---- C ----
    s.label(r, "C. Выбор колонки: наименьшая оставшаяся сумма квадратов   (Теория, раздел 5)", span=5); r += 1
    rc_expl = r; r += 1
    s.header(r, ["Колонка j", "Остаётся Σ(rᵢ − γ)² по обеим группам", "Во сколько раз больше минимума", "Выбор j* = argmin"]); r += 1
    rc = r
    s.cell(r, 2, "Экран", font=Font(name=F, bold=True), fmt=None); s.cell(r, 3, f"=H{rb}+H{rb+1}")
    s.cell(r+1, 2, "SSD", font=Font(name=F, bold=True), fmt=None); s.cell(r+1, 3, f"=H{rb+2}+H{rb+3}")
    s.cell(r, 4, f'=IF(MIN(C{r}:C{r+1})=0,"—",C{r}/MIN(C{r}:C{r+1}))', fmt='0.00"×"'); s.cell(r+1, 4, f'=IF(MIN(C{r}:C{r+1})=0,"—",C{r+1}/MIN(C{r}:C{r+1}))', fmt='0.00"×"')
    ws.merge_cells(start_row=r, start_column=5, end_row=r+1, end_column=5)
    s.cell(r, 5, f'=IF(C{r}<=C{r+1},"экран","SSD")', font=Font(name=F, size=11, bold=True, color=GREEN_F), bg=GREEN_B, fmt=None); s.cell(r+1, 5, None, fmt=None)
    CHOICE = f"$E${rc}"
    ws.conditional_formatting.add(f"B{rc}:D{rc+1}", FormulaRule(formula=[f'$C{rc}=MIN($C${rc}:$C${rc+1})'], fill=fill(GREEN_B), font=Font(name=F, bold=True, color=GREEN_F)))
    s.explain(rc_expl, f'="Складываем ошибку, оставшуюся в двух группах каждой колонки, и берём минимум. Экран: "&{T(f"H{rb}")}&" + "&{T(f"H{rb+1}")}&" = "&{T(f"C{rc}")}&";  SSD: "&{T(f"H{rb+2}")}&" + "&{T(f"H{rb+3}")}&" = "&{T(f"C{rc+1}")}&".  Меньше у: "&{CHOICE}&IF(MIN(C{rc}:C{rc+1})>0," (у соседки в "&TEXT(MAX(C{rc}:C{rc+1})/MIN(C{rc}:C{rc+1}),"0.00")&" раза больше)","")&". Значит h_{m} строится по этой колонке, и каждому объекту достаётся γ его группы."', 44)
    r += 3

    # ---- D ----
    s.label(r, f"D. Обновление модели: F_{m}(xᵢ) = F_{m-1}(xᵢ) + ν·h_{m}(xᵢ)   (Теория, раздел 6)", span=9); r += 1
    rd_expl = r; r += 1
    s.header(r, [c.replace("m−1", str(m-1)).replace("_m", f"_{m}") for c in COLS_D]); r += 1
    t0 = r; rows = [t0+i for i in range(4)]
    for i, (n, _, _, _) in enumerate(data):
        rr = rows[i]
        s.cell(rr, 2, n, font=Font(name=F, bold=True), fmt=None); s.cell(rr, 3, f"={prev_new[i]}")
        s.cell(rr, 4, f'=IF({CHOICE}="экран",{SCR[i]},{SSD[i]})'); s.cell(rr, 5, f"={LRc}*D{rr}")
        s.res(rr, 6, f"=C{rr}+E{rr}"); s.cell(rr, 7, f"={REAL[i]}", fmt=INT); s.res(rr, 8, f"=G{rr}-F{rr}")
        s.cell(rr, 9, f"=ABS(H{rr})"); s.cell(rr, 10, f"=H{rr}^2")
    r = rows[-1] + 1
    s.label(r, "Итого после шага:  Σ|rᵢ|  и  Σrᵢ²", span=7); s.res(r, 9, f"=SUM(I{t0}:I{t0+3})"); s.res(r, 10, f"=SUM(J{t0}:J{t0+3})"); tot = r
    ws.conditional_formatting.add(f"I{t0}:I{t0+3}", DataBarRule(start_type="num", start_value=0, end_type="max", color="A9D18E"))
    s.explain(rd_expl, f'="Прибавляем не всю поправку, а долю ν = "&TEXT({LRc},"0.00")&". A: F_{m}(A) = "&{T(f"C{t0}")}&" + "&TEXT({LRc},"0.00")&" × ("&{T(f"D{t0}")}&") = "&{T(f"F{t0}")}&", остаток "&{T(f"G{t0}")}&" − "&{T(f"F{t0}")}&" = "&{T(f"H{t0}")}&";  D: F_{m}(D) = "&{T(f"C{t0+3}")}&" + "&TEXT({LRc},"0.00")&" × ("&{T(f"D{t0+3}")}&") = "&{T(f"F{t0+3}")}&", остаток "&{T(f"H{t0+3}")}&". "&"Итог: Σ|rᵢ| = "&{T(f"I{tot}")}&" (было "&{T(SA0)}&"), Σrᵢ² = "&{T(f"J{tot}")}&" (было "&{T(SS0)}&"). Невыученная часть вклада колонки «"&{CHOICE}&"» умножилась на (1 − ν) = "&TEXT(1-{LRc},"0.00")&": из "&{T(f"MAX(ABS(D{t0}),ABS(D{t0+1}))")}&" осталось "&{T(f"MAX(ABS(D{t0}),ABS(D{t0+1}))*(1-{LRc})")}&" — эта колонка победит снова, только если остаток её вклада больше невыученного вклада другой колонки (Теория, раздел 7)."', 62)
    # status line
    s.merge(status_row, 2, 13, f'="Выбрана колонка: "&{CHOICE}&"   ·   Σ|rᵢ|: "&{T(SA0)}&" → "&{T(f"I{tot}")}&"   ·   Σrᵢ²: "&{T(SS0)}&" → "&{T(f"J{tot}")}',
            font=Font(name=F, size=10, bold=True, color=NAVY), bg=LIGHT, border=None, align=left, fmt=None)
    step_info.append(dict(m=m, top=top, choice=CHOICE, sa=f"I{tot}", se=f"J{tot}", gE1=f"G{rb}", gE2=f"G{rb+1}", gS1=f"G{rb+2}", gS2=f"G{rb+3}",
                          F=[f"F{rows[i]}" for i in range(4)], amp=f"MAX(ABS(Шаги!E{t0}),ABS(Шаги!E{t0+1}))"))
    prev_new = [f"F{rows[i]}" for i in range(4)]
    r += 2
    # outline group: rows top+1 .. r-2 collapsible under the step banner
    ws.row_dimensions.group(top+1, r-2, outline_level=1, hidden=False)

ws.sheet_properties.outlinePr.summaryBelow = False
for c, w in {1: 3, 2: 13, 3: 20, 4: 18, 5: 16, 6: 16, 7: 18, 8: 18, 9: 15, 10: 13, 11: 10, 12: 10, 13: 10, 14: 2}.items():
    ws.column_dimensions[cl(c)].width = w
ws.freeze_panes = "A6"
ws.print_title_rows = "1:5"; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True

# ============================================================
# Sheet: Сводка
# ============================================================
sv.column_dimensions["A"].width = 2
for c, w in {2: 8, 3: 16, 4: 15, 5: 15, 6: 20, 7: 20, 8: 12, 9: 12, 10: 12, 11: 12, 12: 12}.items(): sv.column_dimensions[cl(c)].width = w
ss.merge(2, 2, 12, "Сводка по 20 шагам и графики", font=Font(name=F, size=15, bold=True, color=NAVY), border=None, align=left, fmt=None)
sv.row_dimensions[2].height = 28
ss.link(3, 2, "← Содержание", "#'Содержание'!A1"); ss.link(3, 4, "Шаги →", "#'Шаги'!A1"); ss.link(3, 6, "Теория →", "#'Теория'!A1")
r = 5
ss.label(r, "Learning rate ν (с листа «Шаги»)", span=3); ss.cell(r, 5, f"={LR}", font=Font(name=F, bold=True), fmt="0.00"); r += 1
ss.label(r, "F₀ = ȳ", span=3); ss.cell(r, 5, f"=Шаги!{MEAN}"); r += 2
ss.header(r, ["Шаг", "Выбрана колонка", "Σ|остатков|", "Σ квадратов", "Поправки экран {13″; 15″}", "Поправки SSD {256; 512}", "F(A)", "F(B)", "F(C)", "F(D)", "|ν·h| шага"]); r += 1
sf = r
ss.cell(r, 2, 0, fmt="0"); ss.cell(r, 3, "— (среднее)", fmt=None)
ss.cell(r, 4, "=" + "+".join(f"ABS(Шаги!{REAL[i]}-Шаги!{MEAN})" for i in range(4)))
ss.cell(r, 5, "=" + "+".join(f"(Шаги!{REAL[i]}-Шаги!{MEAN})^2" for i in range(4)))
ss.cell(r, 6, None, fmt=None); ss.cell(r, 7, None, fmt=None)
for i in range(4): ss.cell(r, 8+i, f"=Шаги!{MEAN}")
ss.cell(r, 12, 0)
r += 1
for info in step_info:
    m = info["m"]
    x = ss.cell(r, 2, m, font=Font(name=F, size=10, color=TEAL, underline="single"), fmt="0"); x.hyperlink = Hyperlink(ref=x.coordinate, location=f"'Шаги'!B{info['top']}", display=str(m))
    ss.cell(r, 3, f"=Шаги!{info['choice']}", font=Font(name=F, size=10, bold=True, color=GREEN_F), bg=GREEN_B, fmt=None)
    ss.cell(r, 4, f"=Шаги!{info['sa']}"); ss.cell(r, 5, f"=Шаги!{info['se']}")
    ss.cell(r, 6, f'="{{"&TEXT(Шаги!{info["gE1"]},"0.00")&"; "&TEXT(Шаги!{info["gE2"]},"0.00")&"}}"', fmt=None)
    ss.cell(r, 7, f'="{{"&TEXT(Шаги!{info["gS1"]},"0.00")&"; "&TEXT(Шаги!{info["gS2"]},"0.00")&"}}"', fmt=None)
    for i in range(4): ss.cell(r, 8+i, f"=Шаги!{info['F'][i]}")
    ss.cell(r, 12, f"={info['amp']}")
    r += 1
sl = r - 1
sv.conditional_formatting.add(f"C{sf+1}:C{sl}", CellIsRule(operator="equal", formula=['"экран"'], fill=fill(LIGHT), font=Font(name=F, bold=True, color=TEAL)))
sv.conditional_formatting.add(f"C{sf+1}:C{sl}", CellIsRule(operator="equal", formula=['"SSD"'], fill=fill(ORANGE_B), font=Font(name=F, bold=True, color=ORANGE_F)))
sv.conditional_formatting.add(f"D{sf}:D{sl}", DataBarRule(start_type="num", start_value=0, end_type="max", color="A9D18E"))
sv.conditional_formatting.add(f"L{sf}:L{sl}", DataBarRule(start_type="num", start_value=0, end_type="max", color="9DC3E6"))
r += 1
ss.explain(r, f'="Читать так: колонка побеждает, пока её невыученный вклад больше вклада соседки. С ν = "&TEXT({LR},"0.00")&" экран держится первые шаги, затем колонки чередуются, и Σ|rᵢ| убывает геометрически — в логарифмической шкале это прямая линия. Цвет в столбце «Выбрана колонка»: синий — экран, оранжевый — SSD; полоски — величина ошибки и размер поправки."', 46)
r += 2
# charts
ch1 = LineChart(); ch1.title = "Σ|остатков| по шагам"; ch1.y_axis.title = "Σ|rᵢ|"; ch1.x_axis.title = "шаг"; ch1.height, ch1.width = 7.5, 15; ch1.legend = None
ch1.add_data(Reference(sv, min_col=4, min_row=sf-1, max_row=sl), titles_from_data=True); ch1.set_categories(Reference(sv, min_col=2, min_row=sf, max_row=sl))
sv.add_chart(ch1, f"B{r}")
ch2 = LineChart(); ch2.title = "Σ|остатков|, логарифмическая шкала (прямая = геометрическое убывание)"; ch2.y_axis.title = "Σ|rᵢ| (log)"; ch2.x_axis.title = "шаг"; ch2.height, ch2.width = 7.5, 15; ch2.legend = None
ch2.add_data(Reference(sv, min_col=4, min_row=sf-1, max_row=sl), titles_from_data=True); ch2.set_categories(Reference(sv, min_col=2, min_row=sf, max_row=sl))
ch2.y_axis.scaling.logBase = 10
sv.add_chart(ch2, f"H{r}")
r += 16
ch3 = LineChart(); ch3.title = "Предсказания F_m(xᵢ) сходятся к реальным ценам"; ch3.y_axis.title = "цена"; ch3.x_axis.title = "шаг"; ch3.height, ch3.width = 7.5, 15
for i in range(4): ch3.add_data(Reference(sv, min_col=8+i, min_row=sf-1, max_row=sl), titles_from_data=True)
ch3.set_categories(Reference(sv, min_col=2, min_row=sf, max_row=sl))
sv.add_chart(ch3, f"B{r}")
ch4 = BarChart(); ch4.type = "col"; ch4.title = "Размер поправки |ν·h_m| на каждом шаге"; ch4.y_axis.title = "|ν·h|"; ch4.x_axis.title = "шаг"; ch4.height, ch4.width = 7.5, 15; ch4.legend = None
ch4.add_data(Reference(sv, min_col=12, min_row=sf-1, max_row=sl), titles_from_data=True); ch4.set_categories(Reference(sv, min_col=2, min_row=sf, max_row=sl))
sv.add_chart(ch4, f"H{r}")
sv.freeze_panes = f"A{sf}"

# ============================================================
# Содержание (TOC) — now that rows are known
# ============================================================
r = TOC_START
s0.banner(r, "Листы пособия", c1=2); r += 2
for text, target, desc in [
    ("Теория", "#'Теория'!A1", "Вывод каждой формулы: почему среднее, почему остаток — антиградиент, критерий выбора колонки, learning rate, сходимость, правила остановки."),
    ("Сводка", "#'Сводка'!A1", "Таблица по 20 шагам с условным форматированием, четыре графика (включая логарифмическую шкалу), ссылки на каждый шаг."),
    ("Шаги", "#'Шаги'!A1", "Параметры (ν, цены) и 20 сворачиваемых блоков: этапы A–D, у каждой таблицы — живое пояснение с формулой и числами."),
]:
    s0.link(r, 2, text, target, font=Font(name=F, size=11, bold=True, color=TEAL, underline="single")); s0.cell(r, 3, desc, font=Font(name=F, size=10), border=None, align=wrap, fmt=None); ws0.row_dimensions[r].height = 30; r += 1
r += 1
s0.banner(r, "Быстрые переходы к шагам", c1=2); r += 2
for info in step_info:
    s0.link(r, 2, f"Шаг {info['m']}", f"#'Шаги'!B{info['top']}")
    s0.cell(r, 3, f'="выбрана колонка: "&Шаги!{info["choice"]}&",  Σ|rᵢ| после шага = "&TEXT(Шаги!{info["sa"]},"0.####")', font=Font(name=F, size=10, color=GREY), border=None, align=left, fmt=None)
    r += 1
r += 1
s0.banner(r, "Как пользоваться", c1=2); r += 2
for t in ["Жёлтые ячейки на листе «Шаги» — единственные входы: цены четырёх ноутбуков и learning rate ν (проверка: 0.01 ≤ ν ≤ 1).",
          "Зелёные ячейки — ключевые результаты этапа. Бежевые строки — пояснения: это формулы, они пересчитываются вместе с таблицами.",
          "Кнопки «−/+» слева от строк на листе «Шаги» сворачивают блок шага; кнопки «1/2» над ними — все блоки сразу.",
          "Поставьте ν = 1 и убедитесь, что получаются ровно два шага исходного примера; ν = 0.1 — и посмотрите, сколько шагов экран держится в лидерах."]:
    s0.cell(r, 2, "•", border=None, fmt=None); s0.cell(r, 3, t, font=Font(name=F, size=10), border=None, align=wrap, fmt=None); ws0.row_dimensions[r].height = 30; r += 1

wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save("boosting_learning_rate.xlsx")
print("ok")
